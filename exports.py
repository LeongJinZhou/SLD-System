"""Fill QIU's retained proposal package; convert that same file for PDF."""
from __future__ import annotations
import copy, io, os, pathlib, re, shutil, subprocess, sys, tempfile, zipfile
from datetime import datetime, timezone, timedelta
from decimal import Decimal, InvalidOperation
ROOT = pathlib.Path(__file__).resolve().parent
VENDOR = ROOT / '.vendor'
if VENDOR.exists(): sys.path.insert(0, str(VENDOR))
KIT_SITE = pathlib.Path(r'C:/Users/QIU/.agency-kit/runtime/site-packages')
if KIT_SITE.exists(): sys.path.insert(0, str(KIT_SITE))
from lxml import etree as ET
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

TEMPLATE = ROOT / 'templates' / 'event-proposal.docx'

def finance_pdf(report, proposal, club):
    """System actual-account report; no original financial template was supplied for this workflow."""
    from html import escape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    font=pathlib.Path('C:/Windows/Fonts/simsun.ttc')
    styles=getSampleStyleSheet()
    if font.exists():
        pdfmetrics.registerFont(TTFont('QIUFinance',str(font),subfontIndex=0))
        report_font='QIUFinance'
    else:
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        report_font='STSong-Light'
    for style in styles.byName.values():style.fontName=report_font
    def p(value,style='BodyText'):return Paragraph(escape(str(value or '—')).replace('\n','<br/>'),styles[style])
    def stamp(value):
        return datetime.fromisoformat(value.replace('Z','+00:00')).astimezone(timezone(timedelta(hours=8))).strftime('%d %b %Y, %H:%M MYT')
    output=io.BytesIO();doc=SimpleDocTemplate(output,pagesize=(595.28,841.89),leftMargin=40,rightMargin=40,topMargin=40,bottomMargin=40,title='Actual event financial report')
    story=[p('QIU · Student Life Division','Heading2'),p('Actual event accounts','Title'),p(club['name'],'Heading2'),p(proposal['title'],'Heading2'),
           p('Status: '+report['status'].replace('_',' ').title()),p('Reference: '+report['id']+' | Revision: '+str(report['revision'])),
           p('Created: '+stamp(report['created_at'])+' | Last recorded: '+stamp(report['updated_at'])),Spacer(1,12)]
    from finance import planned,money
    rows=[[p('MYR'),p('Proposed'),p('Actual'),p('Variance')]]
    for title,key,actualkey in [('Income','revenue','actual_income'),('Expenses','expenses','actual_expenses')]:
        expected=planned(proposal,key);actual=Decimal(report[actualkey]);rows.append([p(title),p(money(expected)),p(f'{actual:.2f}'),p(f'{actual-Decimal(expected)/100:.2f}')])
    summary=Table(rows,colWidths=[155,120,120,120]);summary.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eef2f7')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),9)]))
    story += [summary,Spacer(1,12),p('Transactions and evidence','Heading2')]
    rows=[[p('Type / budget item'),p('Description'),p('Amount'),p('Evidence references')]]
    rows += [[p(r['kind']+' / '+r['budget_ref']),p(r['description']),p('MYR '+r['amount']),p('\n'.join(r['evidence_ids']))] for r in report['rows']]
    table=Table(rows,colWidths=[105,175,80,155],repeatRows=1,splitInRow=1);table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eef2f7')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),10),('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#ccd5df'))]))
    story += [table,Spacer(1,12),p('Variance explanation','Heading2'),p(report['variance_reason']),Spacer(1,12),p('Review history','Heading2')]
    for entry in report.get('history',[]):
        heading=p(entry['action'].title()+' · '+entry['name']);heading.keepWithNext=True
        story.extend([heading,p(entry['actor']+' · '+stamp(entry['timestamp'])),p(entry['comment']),p('Submitted revision '+str(entry['revision'])+' | SHA-256: '+entry['digest']),Spacer(1,10)])
    if not report.get('history'):story.append(p('No SLD review recorded.'))
    story.append(p('Only approved actual transactions affect the club balance. This report does not turn the proposed budget into a transaction.'))
    doc.build(story);return output.getvalue()

def club_pdf(club):
    """Registration report: no original registration template was supplied."""
    import base64
    from html import escape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.enums import TA_LEFT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether
    local_font = pathlib.Path('C:/Windows/Fonts/simsun.ttc')
    if local_font.exists():
        pdfmetrics.registerFont(TTFont('QIUReport',str(local_font),subfontIndex=0))
        report_font = 'QIUReport'
    else:
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        report_font = 'STSong-Light'
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = report_font
        style.alignment = TA_LEFT
    styles['BodyText'].leading = 15
    def p(value, style='BodyText'):
        return Paragraph(escape(str(value or '—')).replace('\n', '<br/>'), styles[style])
    def stamp(value):
        if not value: return '—'
        try:
            return datetime.fromisoformat(value.replace('Z', '+00:00')).astimezone(timezone(timedelta(hours=8))).strftime('%d %b %Y, %H:%M MYT')
        except ValueError: return str(value)
    def picture(value, width, height):
        if not isinstance(value,str) or not re.match(r'^data:image/(png|jpeg|webp);base64,',value): return None
        try:
            img = Image(io.BytesIO(base64.b64decode(value.split(',',1)[1], validate=True)))
            scale = min(width/img.imageWidth, height/img.imageHeight)
            img.drawWidth, img.drawHeight = img.imageWidth*scale, img.imageHeight*scale
            img.hAlign = 'LEFT'
            return img
        except Exception: return None
    stream = io.BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=(595.28,841.89), leftMargin=44, rightMargin=44, topMargin=44, bottomMargin=48,
                            title='Club registration — '+str(club.get('name','')), author='QIU Student Life Division')
    story = [p('QIU · STUDENT LIFE DIVISION','Heading2'), p('Club registration','Title'), p(club.get('name'),'Heading1')]
    logo = picture(club.get('logo'),80,80)
    if logo: story += [logo,Spacer(1,10)]
    fields = [('Status', str(club.get('status','draft')).replace('_',' ').title()), ('Semester',club.get('semester')),
              ('Category',club.get('category')), ('Assigned president',club.get('owner_email')), ('Advisor',club.get('advisor_email')),
              ('Record reference',club.get('id')), ('Revision',club.get('revision')), ('Created',stamp(club.get('created_at'))),
              ('Last saved',stamp(club.get('updated_at')))]
    table = Table([[p(k),p(v)] for k,v in fields], colWidths=[125,382], hAlign='LEFT')
    table.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(0,-1),colors.HexColor('#eef2f7')),
                              ('BOTTOMPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),8)]))
    story += [table,Spacer(1,14),p('Club biography','Heading2'),p(club.get('bio')),Spacer(1,14),p('Committee','Heading2')]
    rows = [[p('Position'),p('Name'),p('University email')]]
    rows += [[p(m.get('position')),p(m.get('name')),p(m.get('email'))] for m in club.get('committee',[])]
    officers = Table(rows,colWidths=[130,145,232], repeatRows=1,hAlign='LEFT')
    officers.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eef2f7')),
                                 ('LINEBELOW',(0,0),(-1,0),.6,colors.HexColor('#ccd5df')),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
    story += [officers,Spacer(1,14)]
    handovers=[entry for entry in club.get('history',[]) if entry.get('action')=='agm_handover']
    if handovers:
        latest=handovers[-1]
        story += [p('Current committee authority','Heading2'),p('The current committee above reflects the SLD-applied AGM handover. Earlier registration signatures refer to their original submitted version.'),
                  p('Approved AGM reference: '+latest.get('report_id','Unavailable')+' | Applied by: '+latest.get('actor','Unavailable')+' | '+stamp(latest.get('timestamp')))]
    story += [p('Review and handover history','Heading2')]
    for entry in club.get('history',[]):
        block = [p(str(entry.get('action','Review')).replace('_',' ').title()+' · '+stamp(entry.get('timestamp') or entry.get('at'))),
                  p(entry.get('name') or entry.get('actor') or entry.get('email')),p(entry.get('comment') or '')]
        if entry.get('report_id'):block.append(p('AGM report reference: '+entry['report_id']))
        if entry.get('revision') is not None:block.append(p('Submitted version: '+str(entry['revision'])))
        if entry.get('digest'):block.append(p('Version SHA-256: '+entry['digest']))
        signature = picture(entry.get('signature'),180,55)
        if signature: block.append(signature)
        block.append(Spacer(1,8))
        story.append(KeepTogether(block))
    if not club.get('history'): story.append(p('No review recorded.'))
    def footer(canvas, document):
        canvas.setFont('Helvetica',8)
        canvas.setFillColor(colors.HexColor('#526177'))
        canvas.drawString(44,29,'QIU Student Life Division | '+('Approved record' if club.get('status')=='approved' else 'Unapproved registration'))
        canvas.drawRightString(551,29,str(document.page))
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return stream.getvalue()
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W = '{' + NS['w'] + '}'

def _text(p, value):
    """Keep paragraph/run formatting; change only the editable text slot."""
    run = p.find(W+'r')
    props = copy.deepcopy(run.find(W+'rPr')) if run is not None and run.find(W+'rPr') is not None else None
    for child in list(p):
        if child.tag != W+'pPr': p.remove(child)
    r = ET.SubElement(p, W+'r')
    if props is not None: r.append(props)
    lines = str(value or '').split('\n')
    if len(lines)>1:
        props_p=p.find(W+'pPr')
        if props_p is None:props_p=ET.SubElement(p,W+'pPr')
        align=props_p.find(W+'jc')
        if align is None:align=ET.SubElement(props_p,W+'jc')
        align.set(W+'val','left')
    for i,line in enumerate(lines):
        if i: ET.SubElement(r, W+'br')
        t=ET.SubElement(r,W+'t'); t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=line

def _cell(tc, value):
    paras=tc.findall(W+'p')
    if not paras: paras=[ET.SubElement(tc,W+'p')]
    _text(paras[0],value)
    for p in paras[1:]: _text(p,'')

def _cells(row): return row.findall(W+'tc')

def _keep_next(p):
    props=p.find(W+'pPr')
    if props is None:props=ET.SubElement(p,W+'pPr')
    keep=props.find(W+'keepNext')
    if keep is None:keep=ET.SubElement(props,W+'keepNext')
    keep.set(W+'val','1')

def _date(value):
    if not value:return ''
    try:
        stamp=datetime.fromisoformat(value.replace('Z','+00:00'))
        if stamp.tzinfo:stamp=stamp.astimezone(timezone(timedelta(hours=8)))
        return stamp.strftime('%d/%m/%Y')
    except ValueError:return str(value)
def _rows(table): return table.findall(W+'tr')
def _field(table, row, cell, value): _cell(_cells(_rows(table)[row])[cell],value)

def _append_block(body,node):
    section=body.find(W+'sectPr')
    if section is not None:section.addprevious(node)
    else:body.append(node)


def _paragraph(value,bold=False,size=20):
    paragraph=ET.Element(W+'p');props=ET.SubElement(paragraph,W+'pPr')
    ET.SubElement(props,W+'spacing',{W+'after':'100'})
    run=ET.SubElement(paragraph,W+'r');formatting=ET.SubElement(run,W+'rPr')
    ET.SubElement(formatting,W+'rFonts',{W+'ascii':'Arial',W+'hAnsi':'Arial'})
    ET.SubElement(formatting,W+'sz',{W+'val':str(size)})
    if bold:ET.SubElement(formatting,W+'b')
    ET.SubElement(run,W+'t').text=str(value or '')
    _text(paragraph,value)
    return paragraph


def _simple_table(headers,rows,widths):
    table=ET.Element(W+'tbl');props=ET.SubElement(table,W+'tblPr')
    ET.SubElement(props,W+'tblW',{W+'w':str(sum(widths)),W+'type':'dxa'})
    borders=ET.SubElement(props,W+'tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:ET.SubElement(borders,W+side,{W+'val':'single',W+'sz':'4',W+'color':'D6DDE5'})
    grid=ET.SubElement(table,W+'tblGrid')
    for width in widths:ET.SubElement(grid,W+'gridCol',{W+'w':str(width)})
    for index,values in enumerate([headers]+rows):
        row=ET.SubElement(table,W+'tr');rp=ET.SubElement(row,W+'trPr');ET.SubElement(rp,W+'cantSplit')
        if index==0:ET.SubElement(rp,W+'tblHeader')
        for value,width in zip(values,widths):
            cell=ET.SubElement(row,W+'tc');cp=ET.SubElement(cell,W+'tcPr');ET.SubElement(cp,W+'tcW',{W+'w':str(width),W+'type':'dxa'})
            if index==0:ET.SubElement(cp,W+'shd',{W+'fill':'EDF1F5'})
            cell.append(_paragraph(value,index==0,18))
    return table


def _stamp(value):
    if not value:return '—'
    return datetime.fromisoformat(value.replace('Z','+00:00')).astimezone(timezone(timedelta(hours=8))).strftime('%d %b %Y\n%H:%M MYT')


def _audit_appendix(body,record):
    if not record.get('submitted_at'):return
    heading=_paragraph('Approval and verification log',True,28)
    ET.SubElement(heading.find(W+'pPr'),W+'pageBreakBefore');_keep_next(heading);_append_block(body,heading)
    _append_block(body,_paragraph('SYSTEM APPROVAL RECORD',False,16))
    _append_block(body,_paragraph('Status: '+record.get('status','').replace('_',' ').title()+'  |  Current submitted version: '+str(record.get('submitted_revision','—'))))
    _append_block(body,_paragraph('Dates and identities below were recorded by the server. Earlier versions remain historical evidence.'))
    rows=[];versions={}
    for entry in record.get('history',[]):
        signer=entry.get('signer') or {};stage=entry.get('stage') or signer.get('slot') or ('preparer' if entry.get('action')=='submitted' else '')
        action={'submitted':'Signed and submitted','approve':'Signed and endorsed' if stage!='sld' else 'Signed and approved','return':'Returned for correction'}.get(entry.get('action'),str(entry.get('action','')).replace('_',' ').title())
        name=signer.get('name') or entry.get('name') or entry.get('actor','Unavailable')
        role=signer.get('designation') or stage.replace('_',' ').title()
        revision=entry.get('revision','—');digest=entry.get('digest','')
        if digest:versions[str(revision)]=digest
        rows.append([_stamp(entry.get('timestamp')),action,name+'\n'+role,'v'+str(revision),entry.get('comment') or '—'])
    if rows:_append_block(body,_simple_table(['Server time','Action','Recorded by','Ver.','Review note'],rows,[1500,1900,2100,700,2900]))
    else:_append_block(body,_paragraph('No review history recorded.'))
    _append_block(body,_paragraph('Version integrity',True,22))
    if record.get('submitted_digest'):versions[str(record['submitted_revision'])]=record['submitted_digest']
    for revision,digest in versions.items():
        _append_block(body,_paragraph('Version '+revision+' · SHA-256\n'+digest[:32]+'\n'+digest[32:],False,16))
    _append_block(body,_paragraph('Only signatures in the assigned current-version cells represent the current submission. A returned version requires renewed signatures.',False,18))

def _place_signature(xml,paragraph,signer,key,archive,rels,added,max_width=2.2,max_height=.6):
    import base64
    from PIL import Image
    wp='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
    a='http://schemas.openxmlformats.org/drawingml/2006/main'
    pic='http://schemas.openxmlformats.org/drawingml/2006/picture'
    r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    relns='http://schemas.openxmlformats.org/package/2006/relationships'
    next_id=max([int(n.get('id','0')) for n in xml.findall('.//{'+wp+'}docPr')]+[0])+1
    data=base64.b64decode(signer['signature'].split(',',1)[1],validate=True)
    with Image.open(io.BytesIO(data)) as im:
        if im.format!='PNG':raise ValueError('Signed proposal contains an invalid signature image.')
        width,height=im.size
    scale=min(max_width*914400/width,max_height*914400/height);cx,cy=int(width*scale),int(height*scale)
    rid='rIdQiuSignature'+key.title()
    if any(n.get('Id')==rid for n in rels):raise ValueError('Signature relationship conflicts with the source template.')
    media='word/media/qiu-signature-'+key+'.png'
    if media in archive.namelist():raise ValueError('Signature media conflicts with the source template.')
    added[media]=data
    ET.SubElement(rels,'{'+relns+'}Relationship',Id=rid,Type=r+'/image',Target='media/'+media.rsplit('/',1)[1])
    _text(paragraph,'')
    for run in paragraph.findall(W+'r'):paragraph.remove(run)
    draw=ET.SubElement(ET.SubElement(paragraph,W+'r'),W+'drawing')
    inline=ET.SubElement(draw,'{'+wp+'}inline',distT='0',distB='0',distL='0',distR='0')
    ET.SubElement(inline,'{'+wp+'}extent',cx=str(cx),cy=str(cy))
    ET.SubElement(inline,'{'+wp+'}docPr',id=str(next_id),name=key+' signature',descr=signer['name']+' — '+key)
    frame=ET.SubElement(inline,'{'+wp+'}cNvGraphicFramePr');ET.SubElement(frame,'{'+a+'}graphicFrameLocks',noChangeAspect='1')
    graphic=ET.SubElement(inline,'{'+a+'}graphic');gd=ET.SubElement(graphic,'{'+a+'}graphicData',uri=pic)
    picture=ET.SubElement(gd,'{'+pic+'}pic');nv=ET.SubElement(picture,'{'+pic+'}nvPicPr')
    ET.SubElement(nv,'{'+pic+'}cNvPr',id='0',name=media.rsplit('/',1)[1]);ET.SubElement(nv,'{'+pic+'}cNvPicPr')
    fill=ET.SubElement(picture,'{'+pic+'}blipFill');ET.SubElement(fill,'{'+a+'}blip',{'{'+r+'}embed':rid})
    ET.SubElement(ET.SubElement(fill,'{'+a+'}stretch'),'{'+a+'}fillRect')
    sp=ET.SubElement(picture,'{'+pic+'}spPr');xf=ET.SubElement(sp,'{'+a+'}xfrm')
    ET.SubElement(xf,'{'+a+'}off',x='0',y='0');ET.SubElement(xf,'{'+a+'}ext',cx=str(cx),cy=str(cy))
    ET.SubElement(ET.SubElement(sp,'{'+a+'}prstGeom',prst='rect'),'{'+a+'}avLst')

def _signature_cells(xml, table, proposal, club, archive):
    """Place only the current immutable attestations in their assigned template cells."""
    import base64
    from PIL import Image
    wp='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
    a='http://schemas.openxmlformats.org/drawingml/2006/main'
    pic='http://schemas.openxmlformats.org/drawingml/2006/picture'
    r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    relns='http://schemas.openxmlformats.org/package/2006/relationships'
    relpath='word/_rels/document.xml.rels'
    rels=ET.fromstring(archive.read(relpath)); added={}
    exemplar=copy.deepcopy(_cells(_rows(table)[0])[1])
    slots=[('preparer',0,0,'Prepared by:'),('advisor',0,1,'Reviewed by:'),('sld',1,0,'Reviewed by:')]
    next_id=max([int(n.get('id','0')) for n in xml.findall('.//{'+wp+'}docPr')]+[0])+1
    for key,row,column,heading in slots:
        cell=_cells(_rows(table)[row])[column]
        if key=='sld':
            for child in list(cell):
                if child.tag!=W+'tcPr':cell.remove(child)
            for para in exemplar.findall(W+'p'):
                cloned=copy.deepcopy(para)
                for attr in list(cloned.attrib):cloned.attrib.pop(attr)
                cell.append(cloned)
        paras=cell.findall(W+'p');_text(paras[0],heading)
        signer=proposal.get('signatures',{}).get(key,{})
        valid=bool(proposal.get('status') not in ('draft','returned') and signer.get('slot')==key
                   and signer.get('revision')==proposal.get('submitted_revision')
                   and signer.get('digest') and signer.get('digest')==proposal.get('submitted_digest'))
        profile=signer if valid else (club.get('advisor_profile',{}) if key=='advisor' else {})
        for para in paras:
            text=''.join(para.itertext()).strip()
            for label,field in [('Name:','name'),('Designation:','designation'),('Department:','department'),('Date:','signed_at')]:
                if text.startswith(label):
                    value=profile.get(field,'') if field!='signed_at' or valid else ''
                    _text(para,label+' '+(_date(value) if field=='signed_at' else value));break
        if not valid:continue
        _place_signature(xml,paras[1],signer,key,archive,rels,added)
        for para in paras[2:5]:cell.remove(para)
    if added:
        added[relpath]=ET.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
        ctpath='[Content_Types].xml';ct=ET.fromstring(archive.read(ctpath));ctns='http://schemas.openxmlformats.org/package/2006/content-types'
        if not any(n.get('Extension')=='png' for n in ct):
            ET.SubElement(ct,'{'+ctns+'}Default',Extension='png',ContentType='image/png')
            added[ctpath]=ET.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
    return added

def money(value):
    try:
        v=Decimal(str(value or 0))
        if not v.is_finite() or v<0 or v>Decimal('100000000'): raise ValueError('Budget amounts must be between 0 and 100,000,000.')
        return v
    except InvalidOperation: raise ValueError('Invalid budget number.')

def total(items): return sum((money(x.get('price'))*money(x.get('quantity')) for x in items),Decimal(0))
def amount(v): return f'{v:,.2f}'

def _repeat(table, items, fields):
    rows=_rows(table); exemplar=copy.deepcopy(rows[1])
    while len(rows)-1 > max(1,len(items)):
        table.remove(rows[-1]);rows=_rows(table)
    while len(rows)-1 < len(items):
        table.append(copy.deepcopy(exemplar));rows=_rows(table)
    header=rows[0].find(W+'trPr')
    if header is None:header=ET.SubElement(rows[0],W+'trPr')
    repeat=header.find(W+'tblHeader')
    if repeat is None:repeat=ET.SubElement(header,W+'tblHeader')
    repeat.set(W+'val','1')
    for i,row in enumerate(rows[1:]):
        item=items[i] if i<len(items) else {}
        for c,k in zip(_cells(row),fields): _cell(c,item.get(k,''))


def _current_report_signer(report,slot):
    signer=report.get('signatures',{}).get(slot,{})
    if report.get('status') in ('draft','returned') or signer.get('slot')!=slot or signer.get('revision')!=report.get('submitted_revision') or not signer.get('digest') or signer.get('digest')!=report.get('submitted_digest'):return None
    return signer


def _report_signature_cell(xml,cell,heading,report,slot,archive,rels,added,max_width=1.65):
    for node in list(cell):
        if node.tag!=W+'tcPr':cell.remove(node)
    cell.append(_paragraph(heading,True))
    image_paragraph=_paragraph('');cell.append(image_paragraph)
    signer=_current_report_signer(report,slot)
    if signer:_place_signature(xml,image_paragraph,signer,slot,archive,rels,added,max_width)
    else:cell.append(_paragraph('____________________'))
    for label,key in [('Name','name'),('Designation','designation'),('Department','department'),('Date','signed_at')]:
        value=signer.get(key,'') if signer else ''
        cell.append(_paragraph(label+': '+(_date(value) if key=='signed_at' else value),False,18))


def _report_package(archive,xml,rels,added):
    if added:
        added['word/_rels/document.xml.rels']=ET.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
        path='[Content_Types].xml';ct=ET.fromstring(archive.read(path))
        if not any(n.get('Extension')=='png' for n in ct):
            ET.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='png',ContentType='image/png')
            added[path]=ET.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
    data=ET.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True);stream=io.BytesIO()
    with zipfile.ZipFile(stream,'w') as result:
        for entry in archive.infolist():result.writestr(entry,data if entry.filename=='word/document.xml' else added.pop(entry.filename,archive.read(entry.filename)))
        for path,value in added.items():result.writestr(path,value)
    return stream.getvalue()


def _report_logo(archive,xml,club,rels,added):
    """Replace only the source Club Logo shape; retain its position and university artwork."""
    import base64
    from PIL import Image
    a='{http://schemas.openxmlformats.org/drawingml/2006/main}'
    paths=['word/document.xml']+[p for p in archive.namelist() if re.fullmatch(r'word/header\d+\.xml',p)]
    for path in paths:
        part=xml if path=='word/document.xml' else ET.fromstring(archive.read(path))
        placeholders=part.xpath('//w:t[text()="Club Logo"]',namespaces=NS)
        if not placeholders:continue
        logo=club.get('logo','')
        if not logo:
            for node in placeholders:node.text=club['name']
            if path!='word/document.xml':added[path]=ET.tostring(part,xml_declaration=True,encoding='UTF-8',standalone=True)
            continue
        with Image.open(io.BytesIO(base64.b64decode(logo.split(',',1)[1],validate=True))) as im:
            output=io.BytesIO();im.convert('RGBA').save(output,format='PNG')
        normalized='data:image/png;base64,'+base64.b64encode(output.getvalue()).decode()
        relpath='word/_rels/'+path.rsplit('/',1)[1]+'.rels'
        part_rels=rels if path=='word/document.xml' else ET.fromstring(archive.read(relpath))
        for node in placeholders:
            drawing=next((parent for parent in node.iterancestors() if parent.tag==W+'drawing'),None)
            if drawing is None:raise ValueError('Club logo source shape changed.')
            temporary=_paragraph('');_place_signature(part,temporary,dict(signature=normalized,name=club['name']),'ClubLogo',archive,part_rels,added,1.4,1.4)
            graphic=drawing.find('.//'+a+'graphicData');new_graphic=temporary.find('.//'+a+'graphicData')
            graphic.getparent().replace(graphic,copy.deepcopy(new_graphic))
        if path!='word/document.xml':
            added[path]=ET.tostring(part,xml_declaration=True,encoding='UTF-8',standalone=True)
            added[relpath]=ET.tostring(part_rels,xml_declaration=True,encoding='UTF-8',standalone=True)


def club_report_docx(report,club):
    """Fill retained AGM/minutes content sections; add the user's four-stage attestation."""
    template=ROOT/'templates'/('agm-report.docx' if report['type']=='agm' else 'meeting-minutes.docx')
    with zipfile.ZipFile(template) as archive:
        xml=ET.fromstring(archive.read('word/document.xml'));body=xml.find(W+'body');paras=body.findall(W+'p');tables=xml.xpath('//w:tbl',namespaces=NS)
        rels=ET.fromstring(archive.read('word/_rels/document.xml.rels'));added={}
        _report_logo(archive,xml,club,rels,added)
        when=report.get('start_time','')+' – '+report.get('end_time','')+' MYT'
        if report['type']=='agm':
            if len(tables)!=4 or len(paras)<67:raise ValueError('AGM source template changed; review its mapping.')
            updates={5:'SESSION '+str(report.get('report_year') or report.get('semester','')),8:'DATE: '+_date(report.get('date')),10:'TIME: '+when,12:'VENUE: '+report.get('venue',''),14:'STUDENT ORGANIZATION: '+club['name'],
                     28:'\n'.join(str(i)+'. '+s for i,s in enumerate(report.get('objectives',[]),1)),29:'',30:'',
                     35:'Recorded members: '+str(report.get('member_count') if report.get('member_count') is not None else len(report.get('attendance',[])))+'; present: '+str(sum(p.get('status')=='present' for p in report.get('attendance',[])))+'; absent in attendance list: '+str(sum(p.get('status')=='absent' for p in report.get('attendance',[])))+'.',
                     42:report.get('opening_speech',''),43:'',47:'',49:'5.0 OUTGOING COMMITTEE MEMBERS',55:'6.0 ELECTION OF NEW COMMITTEE MEMBERS',57:report.get('election_summary',''),60:'7.0 APPOINTMENT OF NEW COMMITTEE MEMBERS',62:'',66:report.get('adjournment','')}
            for index,value in updates.items():_text(paras[index],value)
            # Retained source whitespace otherwise creates a whole blank page after the signed cover.
            for index in range(16,26):
                if paras[index].getparent() is body and not paras[index].findall('.//'+W+'drawing'):body.remove(paras[index])
            for index in [28,29,30]:
                props=paras[index].find(W+'pPr');numbering=props.find(W+'numPr') if props is not None else None
                if numbering is not None:props.remove(numbering)
            for index in [29,30]:body.remove(paras[index])
            paras[27].addnext(_paragraph(report.get('summary','')))
            positions={p.get('email'):p.get('position','Member') for p in report.get('old_committee',[])}
            positions[report.get('old_advisor_email',club.get('advisor_email'))]='Advisor'
            attendance=[dict(p,no=str(i),name=p.get('name','')+'\nAttendance: '+p.get('status','').title(),position=positions.get(p.get('email'),'Member'),year_group=p.get('year_group') or 'Unavailable') for i,p in enumerate(report.get('attendance',[]),1)]
            _repeat(tables[1],attendance,['no','name','position','university_id','email','year_group','programme'])
            _repeat(tables[2],report.get('agenda',[]),['start','end','agenda'])
            def officers(rows):return [dict(p,no=str(i)) for i,p in enumerate(rows,1)]
            _repeat(tables[3],officers(report.get('old_committee',[])),['no','name','position','university_id','email','programme'])
            tables[3].addnext(_paragraph('Outgoing committee accomplishments\n'+report.get('outgoing_activities','')))
            new_table=copy.deepcopy(tables[3]);_repeat(new_table,officers(report.get('new_committee',[])),['no','name','position','university_id','email','programme']);paras[62].addnext(new_table)
            for cell,slot,heading in zip(_cells(_rows(tables[0])[0]),['secretary','president','advisor'],['Prepared by','Checked by','Reviewed by']):
                _report_signature_cell(xml,cell,heading,report,slot,archive,rels,added)
            sld_table=_simple_table(['SLD verification'],[['']],[9100]);_report_signature_cell(xml,_cells(_rows(sld_table)[1])[0],'Reviewed and approved by SLD',report,'sld',archive,rels,added,2.2);_append_block(body,sld_table)
        else:
            if len(tables)!=2 or len(paras)<48:raise ValueError('Meeting source template changed; review its mapping.')
            updates={2:club['name'],4:'DATE: '+_date(report.get('date')),5:'VENUE: '+report.get('venue',''),6:'TIME: '+when,11:'',13:'',14:'',15:'',16:'',17:'',19:'Purpose: '+report.get('summary',''),22:report.get('minutes',''),23:'',24:'',25:'',26:'',27:'',32:'Not applicable' if not report.get('decisions') else '',37:'Date: '+_date(report.get('next_meeting',{}).get('date')),39:'Time: '+report.get('next_meeting',{}).get('time',''),41:'Venue: '+report.get('next_meeting',{}).get('venue','')}
            for index,value in updates.items():_text(paras[index],value)
            # Remove source example-only paragraphs; preserve headings, imagery and action table.
            for index in [11,13,14,15,16,17,23,24,25,26,27]:
                if paras[index].getparent() is body:body.remove(paras[index])
            _repeat(tables[0],report.get('decisions',[]),['task','person','deadline'])
            signature_table=tables[1];sld_row=copy.deepcopy(_rows(signature_table)[2]);signature_table.append(sld_row)
            for row,slot in zip(_rows(signature_table),['secretary','president','advisor','sld']):
                cells=_cells(row);_cell(cells[0],{'secretary':'Secretary signature','president':'President signature','advisor':'Advisor signature','sld':'SLD signature'}[slot])
                _report_signature_cell(xml,cells[2],'Signed by',report,slot,archive,rels,added,2.2)
                if len(cells)>5:_cell(cells[5],_date((_current_report_signer(report,slot) or {}).get('signed_at')))
            if report.get('attendance'):
                _append_block(body,_paragraph('Activity attendance',True,24))
                _append_block(body,_simple_table(['No.','Student','University email','Attendance'],[[str(i),p['name'],p['email'],p['status'].title()] for i,p in enumerate(report['attendance'],1)],[600,2600,4100,1800]))
        for table in xml.xpath('//w:tbl',namespaces=NS):
            for row in _rows(table):
                props=row.find(W+'trPr')
                if props is None:props=ET.SubElement(row,W+'trPr')
                keep=props.find(W+'cantSplit')
                if keep is None:keep=ET.SubElement(props,W+'cantSplit')
                keep.set(W+'val','1')
            for para in _rows(table)[0].findall('.//'+W+'p'):_keep_next(para)
        for p in paras:
            if re.match(r'^\d[. ]', ''.join(p.itertext()).strip()):_keep_next(p)
        _audit_appendix(body,report)
        return _report_package(archive,xml,rels,added)


def club_report_pdf(report,club):
    if report['type']=='agm' or 'sport' not in club.get('category','').lower():return docx_pdf(club_report_docx(report,club))
    return _attendance_report_pdf(report,club)


def _attendance_report_pdf(report,club):
    """Sports report based on supplied attendance columns; no combined four-slot source PDF exists."""
    import base64,json
    from html import escape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,KeepTogether,PageBreak
    styles=getSampleStyleSheet()
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    font='Helvetica'
    if re.search(r'[\u3400-\u9fff]',json.dumps([report,club],ensure_ascii=False)):
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'));font='STSong-Light'
    for style in styles.byName.values():style.fontName=font
    def p(value,style='BodyText'):return Paragraph(escape(str(value or '—')).replace('\n','<br/>'),styles[style])
    def table(rows,widths):
        value=Table([[p(cell) for cell in row] for row in rows],colWidths=widths,repeatRows=1,splitInRow=1)
        value.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#edf1f5')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),9),('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#d6dde5'))]));return value
    stream=io.BytesIO();story=[p('QIU Student Life Division','Heading2'),p('Weekly activity attendance report','Title'),p(club['name'],'Heading2'),p(report['title']),p('Semester: '+report['semester']),p('Activity: '+_date(report['date'])+' · '+report['start_time']+'–'+report['end_time']+' MYT'),p('Venue: '+report['venue']),p('Status: '+report['status'].replace('_',' ').title()),Spacer(1,12),p('Purpose','Heading2'),p(report['summary']),Spacer(1,12)]
    story.append(table([['No.','Student name','University email','Attendance']]+[[str(i),r['name'],r['email'],r['status'].title()] for i,r in enumerate(report.get('attendance',[]),1)],[28,155,242,90]))
    story.extend([Spacer(1,14),p('Approval and verification','Heading2')])
    for slot in ['president','secretary','advisor','sld']:
        signer=_current_report_signer(report,slot);block=[p(slot.title()+' signature','Heading3')]
        if signer:
            im=Image(io.BytesIO(base64.b64decode(signer['signature'].split(',',1)[1])));scale=min(180/im.imageWidth,45/im.imageHeight);im.drawWidth=im.imageWidth*scale;im.drawHeight=im.imageHeight*scale;im.hAlign='LEFT';block.append(im)
            block.append(p(signer['name']+'\n'+signer['designation']+' · '+signer['department']+'\n'+_stamp(signer['signed_at'])))
        else:block.append(p('Awaiting signature'))
        story.append(KeepTogether(block+[Spacer(1,10)]))
    if report.get('submitted_at'):
        story.extend([PageBreak(),p('Approval and verification log','Title'),p('Server-recorded actions for each submitted version.')])
        rows=[['Server time','Action','Recorded by','Version','Review note']]
        for h in report.get('history',[]):
            signer=h.get('signer') or {};rows.append([_stamp(h['timestamp']),h['action'].replace('_',' ').title(),signer.get('name') or h.get('name') or h.get('actor'),str(h.get('revision','')),h.get('comment') or '—'])
        story.append(table(rows,[90,80,115,45,185]));story.extend([Spacer(1,12),p('Current version integrity','Heading2'),p(report.get('submitted_digest'))])
    def footer(canvas,doc):
        canvas.setFont('Helvetica',8);canvas.drawString(40,25,'QIU Student Life Division · '+report['id']);canvas.drawRightString(555,25,str(doc.page))
    SimpleDocTemplate(stream,pagesize=(595.28,841.89),leftMargin=40,rightMargin=40,topMargin=40,bottomMargin=44).build(story,onFirstPage=footer,onLaterPages=footer)
    return stream.getvalue()

def proposal_docx(proposal, club, user):
    from server import derived_duration
    proposal=dict(proposal,**derived_duration(proposal))
    with zipfile.ZipFile(TEMPLATE) as z:
        xml=ET.fromstring(z.read('word/document.xml'))
        body=xml.find(W+'body');paras=body.findall(W+'p');tables=xml.xpath('//w:tbl',namespaces=NS)
        if len(tables)!=6 or len(paras)<68: raise ValueError('Proposal template layout changed; field mapping must be reviewed.')
        p=proposal
        name=club.get('name',''); title=p.get('title','Untitled event')
        dates=_date(p.get('date',''))+((' – '+_date(p['end_date'])) if p.get('end_date') and p['end_date']!=p.get('date') else '')
        times=' – '.join(filter(None,[p.get('start_time'),p.get('end_time')]))
        revenue=p.get('revenue',[]);expenses=p.get('expenses',[])
        rev=total(revenue);exp=total(expenses)
        f=tables[0]
        for row,value in {1:name,2:_date(p.get('submitted_at') or p.get('created_at') or ''),4:title,5:p.get('summary',''),12:p.get('duration',''),13:p.get('venue',''),25:p.get('collaboration',''),26:'RM '+amount(exp)}.items():_field(f,row,1,value)
        _field(f,3,1,dates);_field(f,3,3,times)
        cca=_cells(_rows(f)[11])[2];cp=cca.findall(W+'p')
        if len(cp)>=2:
            _text(cp[0],': '+str(p.get('cca_participants',''))+' CCA Hours');_text(cp[1],': '+str(p.get('cca_committee',''))+' CCA Hours')
        else:_cell(cca,f": {p.get('cca_participants','')} CCA Hours\n: {p.get('cca_committee','')} CCA Hours")
        natures=p.get('nature',[])
        if isinstance(natures,str):natures=[natures]
        categories=[['Sports/Recreational','Competition','Humanity/Volunteerism'],['Social Events/Festivals','Fundraising','Environmental Initiatives'],['Seminars/Conference/Workshop','Exhibition/Convention','Others']]
        for row,cats in zip([7,8,9],categories):
            for cell,cat in zip([2,4,6],cats):_field(f,row,cell,'✓' if cat in natures else '')
        _field(f,10,1,'Please specify: '+(p.get('nature_other','') if 'Others' in natures else ''))
        _field(f,14,2,'✓' if p.get('mode','Physical')=='Physical' else '')
        _field(f,14,4,'✓' if p.get('mode')=='Virtual' else '')
        goals={str(g) for g in p.get('sdgs',[])}
        for row in range(15,25):
            _field(f,row,2,'✓' if str(row-14) in goals else '')
            _field(f,row,4,'✓' if str(row-4) in goals else '')
        funding=p.get('funding_source',[])
        if isinstance(funding,str):funding=[funding]
        for c,label in [(2,'Own budget'),(4,'Club Fund'),(6,'Request budget from QIU')]:_field(f,27,c,'✓' if label in funding else '')
        logistics = p.get('logistics','')+'\n'+p.get('transport','') if p.get('logistics_required',True) else 'Not applicable'
        replacements={9:title.upper()+' PROPOSAL',12:dates,15:p.get('venue',''),18:p.get('duration',''),21:times,24:name,26:p.get('introduction',''),34:'Expected number of participants: '+str(p.get('target_count','')),35:'Eligibility: '+p.get('eligibility',''),42:logistics,43:'Departure Time: '+p.get('departure_time',''),44:'Pickup Point: '+p.get('pickup_point',''),45:'Estimated Return Time: '+p.get('return_time',''),46:'Pickup Point: '+p.get('return_pickup',''),54:p.get('conclusion','')}
        for i,value in replacements.items():_text(paras[i],value)
        for para in paras:
            numbering=para.find('./'+W+'pPr/'+W+'numPr/'+W+'numId')
            if numbering is not None and numbering.get(W+'val')=='1':_keep_next(para)
        objectives=p.get('objectives',[])
        if isinstance(objectives,str): objectives=objectives.splitlines()
        objectives=[re.sub(r'^\s*(?:\d+[.)]|[IVX]+[.)]|[-•])\s*','',line.strip()) for line in objectives if line.strip()]
        exemplar=copy.deepcopy(paras[30]);anchor=paras[30]
        _text(anchor,objectives[0] if objectives else '')
        body.remove(paras[31])
        for objective in objectives[1:]:
            slot=copy.deepcopy(exemplar);_text(slot,objective);anchor.addnext(slot);anchor=slot
        ict = p.get('ict','') if p.get('ict_required',True) else 'Not applicable'
        if ict:
            # The supplied heading ends with a page break. Put content before it.
            for br in paras[40].findall('.//'+W+'br'):
                if br.get(W+'type')=='page':br.getparent().remove(br)
            slot=copy.deepcopy(paras[26]);_text(slot,ict);paras[40].addnext(slot)
            ET.SubElement(ET.SubElement(slot,W+'r'),W+'br').set(W+'type','page')
            _keep_next(paras[40])
        _repeat(tables[1],p.get('activities',[]),['name','description'])
        _repeat(tables[2],p.get('itinerary',[]),['start','end','agenda'])
        _repeat(tables[3],p.get('committee',[]),['name','role','responsibility'])
        # Clone the supplied table's rows, not a new table design.
        b=tables[4];rows=_rows(b); template_row=copy.deepcopy(rows[2]); prepared=[]
        prepared.append(copy.deepcopy(rows[0]))
        for label,items,subtotal,section,totalrow in [('REVENUE',revenue,rev,rows[1],rows[8]),('EXPENSES',expenses,exp,rows[10],rows[15])]:
            prepared.append(copy.deepcopy(section))
            for i,item in enumerate(items or [{}],1):
                row=copy.deepcopy(template_row)
                values=[str(i) if item else '',item.get('description',''),amount(money(item.get('price'))) if item else '',str(item.get('quantity','')),amount(money(item.get('price'))*money(item.get('quantity'))) if item else '']
                for c,value in zip(_cells(row),values):_cell(c,value)
                prepared.append(row)
            row=copy.deepcopy(totalrow);_cell(_cells(row)[-1],amount(subtotal));prepared.append(row)
            prepared.append(copy.deepcopy(rows[9]))
        last=copy.deepcopy(rows[-1]);_cell(_cells(last)[-1],amount(rev-exp));prepared.append(last)
        for row in rows:b.remove(row)
        for row in prepared:b.append(row)
        # Keep each row intact and the closing totals/profit summary together.
        for row in prepared:
            props=row.find(W+'trPr')
            if props is None:props=ET.SubElement(row,W+'trPr')
            if props.find(W+'cantSplit') is None:ET.SubElement(props,W+'cantSplit')
        for row in prepared[-3:-1]:
            for para in row.findall('.//'+W+'p'):_keep_next(para)
        for table in tables[1:5]:
            rows=_rows(table)
            for row in rows:
                props=row.find(W+'trPr')
                if props is None:props=ET.SubElement(row,W+'trPr')
                if props.find(W+'cantSplit') is None:ET.SubElement(props,W+'cantSplit')
            header=rows[0];props=header.find(W+'trPr')
            if props.find(W+'tblHeader') is None:ET.SubElement(props,W+'tblHeader')
            for para in header.findall('.//'+W+'p'):_keep_next(para)
        additions=_signature_cells(xml,tables[5],p,club,z)
        _audit_appendix(body,p)
        data=ET.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
        out=io.BytesIO()
        with zipfile.ZipFile(out,'w') as result:
            for entry in z.infolist():result.writestr(entry,data if entry.filename=='word/document.xml' else additions.pop(entry.filename,z.read(entry.filename)))
            for filename,value in additions.items():result.writestr(filename,value)
        return out.getvalue()

def renderer():
    configured=os.environ.get('SOFFICE_PATH')
    if configured and pathlib.Path(configured).is_file():return configured
    local=ROOT/'.tools'/'libreoffice'/'program'/'soffice.exe'
    if local.is_file():return str(local)
    # Administrative MSI extraction includes a product folder on some versions.
    if (ROOT/'.tools'/'libreoffice').exists():
        found=next((ROOT/'.tools'/'libreoffice').rglob('soffice.exe'),None)
        if found:return str(found)
    return shutil.which('soffice') or shutil.which('soffice.exe')

def pdf_available():return bool(renderer())

def docx_pdf(data):
    engine=renderer()
    if not engine:raise ValueError('PDF renderer is not configured. Word export is available; configure SOFFICE_PATH for template-based PDF.')
    with tempfile.TemporaryDirectory(prefix='qiu-export-') as td:
        folder=pathlib.Path(td);source=folder/'proposal.docx';source.write_bytes(data)
        process=subprocess.run([engine,'-env:UserInstallation='+(folder/'profile').as_uri(),'--headless','--convert-to','pdf','--outdir',str(folder),str(source)],capture_output=True,timeout=90)
        pdf=folder/'proposal.pdf'
        if process.returncode or not pdf.exists():raise ValueError('The original-template PDF could not be rendered. Download Word and retry after checking the renderer.')
        return pdf.read_bytes()

def proposal_xlsx(p,club):
    w=Workbook();s=w.active;s.title='Event budget'
    def row(values):
        s.append(values)
        for c in s[s.max_row]:
            if isinstance(c.value,str):c.data_type='s' # prevent user text becoming a formula
    row(['Event',p.get('title','')]);row(['Club',club.get('name','')]);row(['Date',p.get('date','')]);row([])
    row(['Type','Description','Unit price (RM)','Quantity','Total (RM)'])
    for label,key in [('Revenue','revenue'),('Expense','expenses')]:
        for item in p.get(key,[]):
            row([label,item.get('description',''),float(money(item.get('price'))),float(money(item.get('quantity'))),float(money(item.get('price'))*money(item.get('quantity')))])
    row([]);row(['Total revenue','','','',float(total(p.get('revenue',[])))]);row(['Total expenses','','','',float(total(p.get('expenses',[])))]);row(['Surplus / shortfall','','','',float(total(p.get('revenue',[]))-total(p.get('expenses',[])))])
    for c in s[5]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='183B56')
    for col,width in [('A',22),('B',55),('C',22),('D',15),('E',23)]:s.column_dimensions[col].width=width
    for cells in s.iter_rows(min_row=6):
        for c in cells:
            c.alignment=Alignment(vertical='top',wrap_text=True)
            if c.column in (3,5) and isinstance(c.value,(float,int)):c.number_format='#,##0.00'
    s.freeze_panes='C6';s.auto_filter.ref=f'A5:E{max(5,s.max_row-4)}'
    stream=io.BytesIO();w.save(stream);return stream.getvalue()

def export_proposal(proposal,club,user,format):
    filename=re.sub(r'[^\w\- ]','',proposal.get('title','Event proposal'),flags=re.UNICODE).strip()[:70] or 'Event proposal'
    if format=='xlsx':return proposal_xlsx(proposal,club),filename+'.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    docx=proposal_docx(proposal,club,user)
    if format=='docx':return docx,filename+'.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if format=='pdf':return docx_pdf(docx),filename+'.pdf','application/pdf'
    raise ValueError('Unsupported export format.')
