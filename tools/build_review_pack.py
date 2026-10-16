"""Rebuild the plain-language SLD/ICT review DOCX from its Markdown source."""
import re
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'review'
OUT.mkdir(parents=True, exist_ok=True)
FLOWS = OUT / 'flows'
FLOWS.mkdir(exist_ok=True)
FLOW_STEPS = {
    'registration': (['SLD authorises President\nand advisor', 'President prepares\nclub registration', 'Advisor reviews\nand signs', 'Confirm final send\nto SLD twice', 'SLD opens and locks\nsubmitted registration', 'Approve and open club\nor return for correction'], 'Returned changes need renewed review. Drafts save automatically.'),
    'proposal': (['Committee prepares\nproposal and budget', 'President or Vice President\nsigns one preparation slot', 'Advisor reviews\nand signs', 'Confirm final send\nto SLD twice', 'SLD reviews and signs\nDeficit needs reason', 'Approved event appears\non system calendar'], 'Return with a reason → correct draft → new version and fresh signatures.'),
    'reports': (['Complete weekly or\nannual AGM report', 'Upload required photos\nWeekly 2   AGM 1', 'President signs\nand submits', 'Secretary reviews\nand signs', 'Advisor signs and\ndouble confirms SLD send', 'SLD reviews and signs\nor returns with reason'], 'Actual activities only. A corrected version repeats the required signatures.'),
    'handover': (['Approved AGM with\nincoming committee', 'SLD acknowledges\nand starts handover', 'Incoming committee\ngains pending access', 'Outgoing officers teach\nand keep current titles', 'Server reaches recorded\n7 day default deadline', 'Incoming titles activate\nOutgoing club access ends'], 'SLD controls the overlap days. Independent event assignments remain separate.'),
    'withdrawal': (['Treasurer prepares\nand signs request', 'President reviews\nand signs', 'Advisor signs and\ndouble confirms SLD send', 'SLD reviews\nand signs', 'SLD uploads Finance\nsigned approval', 'SLD all approval\nWaiting for transfer', 'Treasurer uploads receipt\nand signs received', 'SLD acknowledges\nreceipt and closes'], 'Transfer target is 14 days by default. SLD alone changes policy. No duplicate expense posting.'),
}


def font(size, bold=False):
    candidate = Path('C:/Windows/Fonts/' + ('arialbd.ttf' if bold else 'arial.ttf'))
    return ImageFont.truetype(str(candidate), size) if candidate.exists() else ImageFont.load_default(size=size)


def flow(name, steps, note):
    width, height = 1280, 370
    img = Image.new('RGB', (width, height), 'white')
    draw = ImageDraw.Draw(img)
    columns = 4 if len(steps) == 8 else 3
    boxw = (width - 100 - (columns - 1) * 35) // columns
    positions = [(50 + i * (boxw + 35), 30) for i in range(columns)]
    positions += [(50 + i * (boxw + 35), 190) for i in reversed(range(columns))]
    for index, text in enumerate(steps):
        x, y = positions[index]
        draw.rounded_rectangle((x, y, x + boxw, y + 100), radius=9, fill='#eef3f7', outline='#708091', width=2)
        draw.text((x + 14, y + 10), str(index + 1), font=font(19, True), fill='#192b3d')
        draw.multiline_text((x + boxw / 2, y + 61), text, font=font(22), fill='#101820', anchor='mm', align='center', spacing=6)
        if index:
            px, py = positions[index - 1]
            if py == y:
                right = x > px
                start = (px + boxw + 3 if right else px - 3, py + 50)
                end = (x - 4 if right else x + boxw + 4, y + 50)
                draw.line((start, end), fill='#44596b', width=3)
                direction = 1 if right else -1
                draw.polygon([end, (end[0] - direction * 10, end[1] - 6), (end[0] - direction * 10, end[1] + 6)], fill='#44596b')
            else:
                cx = x + boxw / 2
                draw.line(((cx, py + 103), (cx, y - 5)), fill='#44596b', width=3)
                draw.polygon([(cx, y - 4), (cx - 6, y - 14), (cx + 6, y - 14)], fill='#44596b')
    draw.text((width / 2, 333), note, font=font(19), fill='#42515d', anchor='mm')
    img.save(FLOWS / (name + '.png'))


def rich(paragraph, text):
    for i, part in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        run = paragraph.add_run(part)
        run.bold = bool(i % 2)


def table(doc, lines):
    values = [[part.strip() for part in line.strip().strip('|').split('|')] for line in lines if not re.fullmatch(r'[|\s:-]+', line)]
    t = doc.add_table(rows=0, cols=len(values[0]))
    t.style = 'Table Grid'
    for index, values_row in enumerate(values):
        cells = t.add_row().cells
        props = t.rows[-1]._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        props.append(no_split)
        if index == 0:
            repeat = OxmlElement('w:tblHeader')
            props.append(repeat)
        for cell, value in zip(cells, values_row):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.space_before = Pt(4)
            rich(p, value)
            for run in p.runs:
                run.font.size = Pt(9.5)
                if index == 0:
                    run.bold = True
            if index == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E8EDF2')
                cell._tc.get_or_add_tcPr().append(shade)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    for name, (steps, note) in FLOW_STEPS.items():
        flow(name, steps, note)
    doc = Document()
    for border in list(doc.styles.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    section.header_distance = section.footer_distance = Inches(.25)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.styles['Normal'].paragraph_format.space_after = Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.1
    doc.styles['Title'].font.size = Pt(27)
    doc.styles['Heading 1'].font.size = Pt(19)
    doc.styles['Heading 2'].font.size = Pt(14)
    doc.styles['Heading 3'].font.size = Pt(11.5)
    header = section.header.paragraphs[0]
    header.text = 'QIU Student Life Division  |  Requirements review draft'
    header.style = 'Normal'
    header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.add_run('SLD and ICT review  |  9 October 2026                                      Page ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
    lines = (OUT / 'QIU_SLD_Requirements_and_Workflows.md').read_text(encoding='utf-8-sig').splitlines()
    i, cover, code = 0, True, False
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line == '<!-- pagebreak -->':
            doc.add_page_break()
            cover = False
        elif line.startswith('!['):
            path = re.search(r'\]\((.+)\)', line)[1]
            doc.add_picture(str(OUT / path), width=Inches(6.8))
            cap = doc.add_paragraph(re.search(r'!\[(.+)\]', line)[1])
            cap.runs[0].italic = True
            cap.runs[0].font.size = Pt(9)
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                rows.append(lines[i])
                i += 1
            table(doc, rows)
            continue
        elif line.startswith('```'):
            code = not code
        elif code:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(1)
            p.runs[0].font.name = 'Consolas'
            p.runs[0].font.size = Pt(9)
        elif line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title' if cover else 'Heading 1')
        elif line.startswith('## '):
            doc.add_paragraph(line[3:], 'Subtitle' if cover else 'Heading 2')
        elif line.startswith('### '):
            doc.add_paragraph(line[4:], 'Heading 2')
        else:
            style = 'List Bullet' if line.startswith('- ') else 'List Paragraph' if re.match(r'^\d+\. ', line) else None
            p = doc.add_paragraph(style=style)
            rich(p, line[2:] if line.startswith('- ') else line)
        i += 1
    doc.core_properties.title = 'QIU Student Life Division System Requirements and Workflows'
    doc.core_properties.subject = 'Plain language SLD review and ICT handover'
    doc.core_properties.author = 'Prepared for QIU SLD and ICT'
    doc.save(OUT / 'QIU_SLD_Requirements_and_Workflows.docx')
    print('Created review DOCX and five workflow diagrams')


if __name__ == '__main__':
    build()
